from pptx import Presentation

from enums.keys import Key


class ResultCreator():

    def __init__(self, name):
        self.pres = Presentation()
        self.file_name = name + '.pptx'

    def add_slide(self, slide_data: dict):
        slide_type = 1 if Key.SLIDE_TYPE not in slide_data else slide_data[Key.SLIDE_TYPE]
        slide = self.pres.slide_layouts[slide_type]
        page = self.pres.slides.add_slide(slide)
        for key, text in slide_data.items():
            if key == 'title':
                page.shapes.title.text = text
            elif key == 'content' or key == 'left_side_content':
                page.placeholders[1].text = text

            elif key == 'right_side_content':
                page.placeholders[2].text = text


        self.pres.save(self.file_name)
